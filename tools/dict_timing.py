#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""分步计时，找出转换到底卡在哪一步。每步都 flush，避免输出被缓冲看不见。"""
from __future__ import annotations

import re
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "dict" / "洄文大字典.md"

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dict_to_docx as C  # noqa: E402


def t(label: str, fn):
    t0 = time.time()
    r = fn()
    dt = time.time() - t0
    print(f"  {label:<46} {dt:8.2f} 秒", flush=True)
    return r, dt


print("=== 分步计时 ===", flush=True)

lines, _ = t("读文件 + splitlines", lambda: MD.read_text(encoding="utf-8").splitlines())
print(f"    → {len(lines):,} 行", flush=True)

body, _ = t("挑出正文行", lambda: [
    l for l in lines if l.strip() and not l.startswith(("|", "#", ">")) and not l.startswith("```")])
print(f"    → {len(body):,} 行正文", flush=True)

lens = sorted((len(l) for l in body), reverse=True)
print(f"    最长行 {lens[0]:,} 字，中位 {lens[len(lens)//2]} 字，超 1000 字的有 "
      f"{sum(1 for x in lens if x > 1000):,} 行", flush=True)

worst = max(body, key=len)
_, dt_worst = t(f"对最长的一行跑 TOKEN.split（{len(worst):,} 字）",
                lambda: C.TOKEN.split(worst))
if dt_worst > 1:
    print("    ★ 正则回溯爆炸！问题在这里", flush=True)
    print(f"    该行前 200 字：{worst[:200]}", flush=True)

sample = body[:3000]
_, dt3 = t("前 3000 行 TOKEN.split", lambda: [C.TOKEN.split(l) for l in sample])
print(f"    → 外推全量 {dt3 * len(body) / 3000:.1f} 秒", flush=True)

_, dtb = t("前 3000 行 粗体 findall", lambda: [re.findall(r"\*\*.+?\*\*", l) for l in sample])
print(f"    → 外推全量 {dtb * len(body) / 3000:.1f} 秒", flush=True)

# python-docx 本身的开销
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402


def mk(n: int):
    d = Document()
    for i in range(n):
        p = d.add_paragraph()
        C.set_font(p.add_run("测试文本 abc 123"))
    return d


_, dtp = t("新建 3000 个段落各带 1 个 run", lambda: mk(3000))
print(f"    → 外推 80,000 段约 {dtp * 80000 / 3000:.0f} 秒", flush=True)
