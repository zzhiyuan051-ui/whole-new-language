#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""量表格规模，并单独给"造表格"这一步计时。

段落路径已经排除了（外推 9 秒）。剩下最可疑的就是表格：
6936 行、40 块，但**列数**没量过 —— 如果某张表有几十列，
单元格数就是几万，逐个设格式会很慢。
"""
from __future__ import annotations

import re
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "dict" / "洄文大字典.md"

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dict_to_docx as C  # noqa: E402

lines = MD.read_text(encoding="utf-8").splitlines()

# 切出所有表格块
blocks: list[list[list[str]]] = []
i = 0
while i < len(lines):
    if lines[i].startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not re.match(r"^[\s:\-]+$", "".join(cells)):
                rows.append(cells)
            i += 1
        blocks.append(rows)
    else:
        i += 1

print(f"表格块 {len(blocks)} 张", flush=True)
cols = [max((len(r) for r in b), default=0) for b in blocks]
rows_n = [len(b) for b in blocks]
cells = sum(c * r for c, r in zip(cols, rows_n))
print(f"  行数：最多 {max(rows_n)}  合计 {sum(rows_n):,}", flush=True)
print(f"  列数：最多 {max(cols)}  分布 {sorted(set(cols))}", flush=True)
print(f"  ★ 单元格总数 {cells:,}", flush=True)

print("\n=== 给最大的一张表计时 ===", flush=True)
big_i = max(range(len(blocks)), key=lambda k: len(blocks[k]) * cols[k])
b, nc = blocks[big_i], cols[big_i]
print(f"  第 {big_i} 张：{len(b)} 行 × {nc} 列 = {len(b)*nc:,} 单元格", flush=True)

from docx import Document  # noqa: E402

doc = Document()
t0 = time.time()
t = doc.add_table(rows=0, cols=nc)
t.style = "Table Grid"
print(f"  add_table + 设样式        {time.time()-t0:.3f} 秒", flush=True)

t0 = time.time()
for r in b:
    t.add_row()
print(f"  add_row × {len(b):<4}          {time.time()-t0:.3f} 秒", flush=True)

t0 = time.time()
for ri, r in enumerate(b[:50]):
    cells_ = t.rows[ri].cells
    for ci in range(nc):
        cell = cells_[ci]
        cell.text = ""
        p = cell.paragraphs[0]
        C.add_inline(p, r[ci] if ci < len(r) else "", size=9.5)
print(f"  填 50 行（{50*nc:,} 单元格）      {time.time()-t0:.3f} 秒"
      f"  → 外推整张 {len(b)} 行约 {(time.time()-t0)*len(b)/50:.1f} 秒", flush=True)
