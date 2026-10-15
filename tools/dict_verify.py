#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""验证新 docx：Markdown 残留有没有清干净、结构对不对。

判据要硬：
  · 正文里不该再出现 `**`（旧版有 52,760 处）
  · 不该再出现以 `|` 开头的行（旧版把 6,936 行表格原样留成文本）
  · 表格必须是**真 Word 表格**（<w:tbl>），不是文本
  · 标题必须是真标题样式
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

DICT = Path(__file__).resolve().parent.parent / "dict"
NEW = DICT / "洄文大字典.docx"
OLD = DICT / "_旧版备份.docx"


def inspect(p: Path, label: str) -> dict:
    z = zipfile.ZipFile(p)
    doc = z.read("word/document.xml").decode("utf-8")
    text = re.sub(r"<[^>]+>", "", doc)
    paras = len(re.findall(r"<w:p[ >]", doc))
    tbls = len(re.findall(r"<w:tbl>", doc))
    rows = len(re.findall(r"<w:tr[ >]", doc))
    cells = len(re.findall(r"<w:tc>", doc))
    h1 = len(re.findall(r'w:pStyle w:val="Heading1"', doc))
    h2 = len(re.findall(r'w:pStyle w:val="Heading2"', doc))
    h3 = len(re.findall(r'w:pStyle w:val="Heading3"', doc))
    stars = len(re.findall(r"\*\*", text))
    piped = len([l for l in text.split("\n") if l.lstrip().startswith("|")])
    hr = text.count("|---")
    return {
        "文件": p.name, "大小KB": round(p.stat().st_size / 1024),
        "纯文本字数": len(text), "段落": paras, "真表格": tbls,
        "表格行": rows, "单元格": cells,
        "H1": h1, "H2": h2, "H3": h3,
        "残留**": stars, "残留|行": piped, "残留|---": hr,
    }


def residue_outside_code(p: Path) -> tuple[int, list[str]]:
    """只数**代码块之外**的 `**`。

    代码块里的 `**` 是要展示的内容本身（体例里那段条目结构模板就是），
    不该被当成转换残留。靠段落样式名来区分，而不是靠猜。
    """
    from docx import Document
    doc = Document(str(p))
    bad: list[str] = []
    for para in doc.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        if style == "洄文代码块":
            continue
        if "**" in para.text:
            bad.append(para.text[:100])
    # 表格单元格里也查一遍
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for para in c.paragraphs:
                    if "**" in para.text:
                        bad.append("[表格] " + para.text[:100])
    return len(bad), bad[:10]


def show(label: str, d: dict) -> None:
    print(f"=== {label} ({d['文件']}, {d['大小KB']} KB) ===")
    for k in ("纯文本字数", "段落", "真表格", "表格行", "单元格", "H1", "H2", "H3",
              "残留**", "残留|行", "残留|---"):
        print(f"  {k:<12} {d[k]:>9,}")
    print()


def main() -> int:
    new = inspect(NEW, "新版")
    show("新版", new)
    if OLD.exists():
        show("旧版（对照）", inspect(OLD, "旧版"))

    print("=== 逐段检查：代码块之外的 ** 残留 ===")
    n_bad, samples = residue_outside_code(NEW)
    print(f"  {n_bad} 处")
    for s in samples:
        print(f"    {s}")

    print("\n=== 判据 ===")
    checks = [
        ("代码块之外完全没有 ** 残留", n_bad == 0),
        ("完全没有 | 开头残留行", new["残留|行"] == 0),
        ("完全没有 |--- 残留", new["残留|---"] == 0),
        ("表格是真 Word 表格", new["真表格"] >= 40),
        ("标题用了真样式", new["H1"] >= 8 and new["H2"] >= 35 and new["H3"] >= 3000),
    ]
    ok = True
    for name, passed in checks:
        print(f"  {'✔' if passed else '✘'} {name}")
        ok &= passed
    print("\n" + ("全部通过" if ok else "有未通过项"))
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
